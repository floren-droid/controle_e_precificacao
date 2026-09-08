'''
Engine de precificação - Pricing BRAP

Loop interativo de decisão de preço, com:
- Gravação imediata de cada decisão (retomável a qualquer momento - se
  parar no meio, o que já foi decidido está salvo no controle).
- Possibilidade de pular uma pendência e decidir depois.
- Possibilidade de revisar/editar preços já decididos na mesma sessão.
- Ponto de extensão para uma futura regra de preço automática (está
  retornando None por hora enquanto não temos a política de Pricing)
'''

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

STATUS_PRECIFICADO = 'PRECIFICADO'

# ======================================================================
# Ponto de extensão - regra de preço automático (ainda sem uso)

def sugerir_preco(linha:pd.Series) -> float | None:
    '''
    Ainda sem utilidade, mas servirá para devolver um valor sugerido
    automaticamente, dentro das normas da política de preço.
    '''
    return None

# ======================================================================
# 1. Exibição de cada pendência

def exibir_linhas(linha: pd.Series, posicao: int | None = None,
                  total:int | None = None, decididas: int | None = None) -> None:
    if posicao is not None and total is not None:
        extra = f'  ({decididas} já precificada(s))' if decididas is not None else ''
        print(f'Pendências {posicao} de {total}{extra}')
    print('-' * 64)
    print(f'[{int(linha["_indice_controle"])}] {linha["LABORATÓRIO"]} | '
          f'{linha["REPRESENTANTE"]}\n')
    print(f'  Exame: {linha["MN_SHIFT"]} - {linha.get("DESCRIÇÃO_EXAME", "")}')

    custo = linha.get('CUSTO_TOTAL')
    print(f'  Custo total: {custo:.2f}' if pd.notna(custo) else ' Custo total: \
          (sem dado)')

    tabela = linha.get('TABELA_PADRÃO')
    concorrente = linha.get('VALOR_CONCORRENTE')
    negociado = linha.get('NEGOCIADO')
    print(f'  Tabela padrão: {tabela if pd.notna(tabela) else "-"}'
          f'   |   Concorrente: {concorrente if pd.notna(concorrente) else "-"}'
          f'   |   Negociado: {negociado if pd.notna(negociado) else "-"}\n'
    )

    aprovado_atual = pd.to_numeric(linha.get('APROVADO'), errors='coerce') # inclusão de tratar vazio

    if pd.notna(aprovado_atual):
        margem = linha.get('MARGEM')
        texto_margem = f' (margem {margem:.1%})' if pd.notna(margem) else ''
        print(f'  >> Já decidido nesta sessão: {aprovado_atual:.2f}{texto_margem}\n')

# ======================================================================
# 2. Validação de entrada

def _parse_preco(texto: str) -> float | None:
    texto = texto.strip().replace(',', '.')
    if not texto:
        return None
    try:
        valor = float(texto)
    except ValueError:
        return None
    return valor if valor > 0 else None

def _calcular_margem(preco: float, custo: float | None) -> float | None:
    if custo is None or pd.isna(custo) or preco == 0:
        return None
    return (preco - custo) / preco

def _confirmar(mensagem: str = 'Confirmar? (Enter confirm, "n" cancela): ') -> bool:
    resposta = input(mensagem).strip().lower()
    return resposta != "n"

# ======================================================================
# 3. Gravação imeditada no controle

def _mapa_colunas(ws) -> dict:
    return {cell.value: cell.column for cell in ws[1] if cell.value}

def registrar_decisao(wb: Workbook, aba: str, indice_controle: int, preco:float, status: str) \
    -> None:
    '''
    Grava o preço aprovado e o status diretamente na linha correta do
    controle, usando o índice original preservado desde a coleta
    '''
    ws = wb[aba]
    colunas = _mapa_colunas(ws)
    linha_planilha = indice_controle + 2 # +1 cabeçalho e +1 porque Excel começa em 1

    ws.cell(row=linha_planilha, column=colunas['APROVADO'], value=preco)
    ws.cell(row=linha_planilha, column=colunas['STATUS'], value=status)

# ======================================================================
# 4. Revisao / Edição de decisões já tomadas na sessão (refazer preço)

def _revisar_decididos(
        arquivo_analise: pd.DataFrame, wb: Workbook, aba: str, caminho_controle: Path
) -> pd.DataFrame:
      decididos = arquivo_analise[arquivo_analise['STATUS'] == STATUS_PRECIFICADO]

      if decididos.empty:
          print('Nada precificado ainda nesta sessão.')
          input('Enter para voltar...')
          return arquivo_analise

      _limpar_tela()
      print('Precificados nesta sessão:')
      posicoes = list(decididos.index)
      for pos_exibida, idx in enumerate(posicoes):
          linha = arquivo_analise.loc[idx]
          margem = linha.get('MARGEM')
          texto_margem = f'  |  margem{margem:.1%}' if pd.notna(margem) else ''
          print(f'  {pos_exibida}) [{int(linha["_indice_controle"])}] {linha["LABORATÓRIO"]} '
                f'- {linha["MN_SHIFT"]} = {linha["APROVADO"]:.2f}{texto_margem}')

      escolha = input('Número para editar (Enter para voltar sem editar): ').strip()
      if not escolha:
          return arquivo_analise

      try:
          idx_editar = posicoes[int(escolha)]
      except (ValueError, IndexError):
          print('Opções inválida, voltande sem editar.')
          input('Enter para continuar...')
          return arquivo_analise

      linha = arquivo_analise.loc[idx_editar]
      exibir_linhas(linha)
      novo_texto = input('Novo preço aprovado: ').strip()
      novo_preco = _parse_preco(novo_texto)

      if novo_preco is None:
          print('Valor inválida, mantendo o preço anterior.')
          input('Enter para continuar...')
          return arquivo_analise

      novo_margem = _calcular_margem(novo_preco, linha.get('CUSTO_TOTAL'))
      print(f'Novo preço: {novo_preco:.2f}  |  margem {novo_margem:.1%}' if novo_margem is not None
            else f'Novo preço: {novo_preco:.2f}')

      if not _confirmar():
          print('Edição cancelada.')
          input('Enter para continuar...')
          return arquivo_analise
      
      registrar_decisao(wb, aba, int(linha['_indice_controle']), novo_preco, STATUS_PRECIFICADO)
      wb.save(caminho_controle)

      arquivo_analise.loc[idx_editar, 'APROVADO'] = novo_preco
      arquivo_analise.loc[idx_editar, 'MARGEM'] = novo_margem
      print(f'Atualizado para {novo_preco:.2f}.')
      return arquivo_analise

# ======================================================================
# 5. Loop principal

def precificar_interativo(
        arquivo_analise: pd.DataFrame,
        caminho_controle: Path,
        aba: str,
) -> pd.DataFrame:
    '''
    Percorre as pendências permitindo, acada uma:
    - digitar um preço        -> grava na hora, marca PRECIFICADO
    - "p" pular               -> mantém PENDENTE, decide depois
    - "r" revisar/editar já feitos nesta sessão
    - "q" sair a qualquer momento (o que já foi decidido está salvo)

    Retorna o arquivo_analise atualizado com as decisões desta sessão.
    Linhas puladas continuam sem preço - rode de novo quando quiser
    retomar exatamento de onde parou (eles continuam PENDENTE no
    controle).
    '''
    wb = load_workbook(caminho_controle)
    arquivo_analise = arquivo_analise.copy()
    total = len(arquivo_analise)

    i = 0
    while i < len(arquivo_analise):
        idx = arquivo_analise.index[i]
        linha = arquivo_analise.loc[idx]

        decididas = int((arquivo_analise['STATUS'] == STATUS_PRECIFICADO).sum())

        _limpar_tela()
        exibir_linhas(linha, posicao=i + 1, total=total, decididas=decididas)

        sugestao = sugerir_preco(linha)
        dica_sugestao = f' [sugestão: {sugestao:.2f}]' if sugestao is not None else ''
        entrada = input(
            f'Preço aprovado{dica_sugestao} ("p" pular, "r" revisar, "q" sair): '
        ).strip().lower()

        if entrada == 'q':
            print('Saindo. O que foi precificado está salvo no controle.')
            break

        if entrada == 'p':
            i += 1
            continue

        if entrada == 'r':
            arquivo_analise = _revisar_decididos(arquivo_analise, wb, aba, caminho_controle)
            continue # não avança, apenas permite revisar quantas vezes quiser

        preco = _parse_preco(entrada)
        if preco is None:
            print('Valor não válido! Digite um número, "p", "r" ou "q".')
            input('Enter para continuar...')
            continue

        margem = _calcular_margem(preco, linha.get('CUSTO_TOTAL'))
        texto_margem = f'  |  margem resultante: {margem:.1%}' if margem is not None else ''
        print(f'Preço: {preco:.2f}{texto_margem}')

        if not _confirmar():
            print('Preço não confirmado, digite de novo.')
            input('Enter para continuar...')
            continue # não avança, deixa tentar de novo na mesma linha

        registrar_decisao(wb, aba, int(linha['_indice_controle']), preco, STATUS_PRECIFICADO)
        wb.save(caminho_controle)

        arquivo_analise.loc[idx, 'APROVADO'] = preco
        arquivo_analise.loc[idx, 'MARGEM'] = margem
        arquivo_analise.loc[idx, 'STATUS'] = STATUS_PRECIFICADO

        i += 1

    return arquivo_analise

# ======================================================================
# 6. Limpar tela

def _limpar_tela() -> None:
    '''
    Limpa a saída anterior antes de mostrar a próxima pendência.
    '''
    try:
        from IPython.display import clear_output
        clear_output(wait=True)
    except ImportError:
        import os
        os.system('cls' if os.name == 'nt' else 'clear')
